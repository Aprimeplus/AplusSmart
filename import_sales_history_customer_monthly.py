# -*- coding: utf-8 -*-
"""Import ข้อมูลยอดขายรายลูกค้าย้อนหลัง (ปี 2567/2568) จากไฟล์ "SO Analysis Report"
เข้าตาราง sales_history_customer_monthly — รวมยอดต่อ (ลูกค้า, รหัสพนักงานเดิม, ปี, เดือน)
จาก 14,641 แถวระดับรายการสินค้า ให้เหลือแค่ยอดรวมต่อเดือนต่อลูกค้า (ตามที่ Customer Monitoring ใช้)

รันซ้ำได้ปลอดภัย — ใช้ ON CONFLICT อัปเดตทับด้วยยอดล่าสุดเสมอ (ไม่บวกซ้ำ)
"""
import sys
import openpyxl
import pandas as pd
from main_app import AppContainer

SRC_PATH = r"C:\Users\Nitro V15\Downloads\2.SO Analysis Report 2026(2).xlsx"
SHEETS = ["SO All24", "SO All25"]

COL_DATE = 1
COL_CUSTOMER_CODE = 5
COL_CUSTOMER_NAME = 6
COL_SALE_KEY = 7
COL_AMOUNT = 20  # "จำนวนเงินหลังส่วนลดรวม"


def read_sheet(wb, sheet_name):
    ws = wb[sheet_name]
    rows = []
    for r in range(2, ws.max_row + 1):
        date_val = ws.cell(row=r, column=COL_DATE).value
        cust_code = ws.cell(row=r, column=COL_CUSTOMER_CODE).value
        if date_val is None or not cust_code:
            continue
        cust_name = ws.cell(row=r, column=COL_CUSTOMER_NAME).value
        sale_key = ws.cell(row=r, column=COL_SALE_KEY).value
        amount = ws.cell(row=r, column=COL_AMOUNT).value
        try:
            amount = float(amount) if amount not in (None, "") else 0.0
        except (TypeError, ValueError):
            amount = 0.0
        rows.append({
            "customer_code": str(cust_code).strip(),
            "customer_name": (cust_name or "").strip() if isinstance(cust_name, str) else cust_name,
            "sale_key_raw": (sale_key or "").strip() if isinstance(sale_key, str) else sale_key,
            "year": date_val.year,
            "month": date_val.month,
            "amount": amount,
            "source_sheet": sheet_name,
        })
    return rows


def main():
    print(f"Reading {SRC_PATH} ...")
    wb = openpyxl.load_workbook(SRC_PATH, data_only=True)

    all_rows = []
    for sheet in SHEETS:
        rows = read_sheet(wb, sheet)
        print(f"- {sheet}: {len(rows)} รายการสินค้า")
        all_rows.extend(rows)

    df = pd.DataFrame(all_rows)
    print(f"รวมทั้งหมด: {len(df)} รายการสินค้า จาก {len(SHEETS)} sheet")

    # รวมยอดต่อ (ลูกค้า, รหัสพนักงานเดิม, ปี, เดือน) — ใช้ชื่อลูกค้าล่าสุด (แถวท้ายสุดตามลำดับไฟล์) เผื่อชื่อเปลี่ยนระหว่างปี
    grouped = (
        df.groupby(["customer_code", "sale_key_raw", "year", "month"], dropna=False)
        .agg(total_amount=("amount", "sum"), customer_name=("customer_name", "last"), source_sheet=("source_sheet", "last"))
        .reset_index()
    )
    print(f"รวมยอดแล้วเหลือ: {len(grouped)} แถว (ลูกค้า x รหัสพนักงาน x เดือน x ปี)")

    app = AppContainer()
    conn = app.get_connection()
    try:
        with conn.cursor() as cur:
            n_inserted = 0
            for _, r in grouped.iterrows():
                cur.execute("""
                    INSERT INTO sales_history_customer_monthly
                        (customer_code, customer_name, sale_key_raw, year, month, total_amount, source_sheet)
                    VALUES (%s, %s, %s, %s, %s, %s, %s)
                    ON CONFLICT (customer_code, sale_key_raw, year, month) DO UPDATE SET
                        customer_name = EXCLUDED.customer_name,
                        total_amount = EXCLUDED.total_amount,
                        source_sheet = EXCLUDED.source_sheet,
                        imported_at = NOW()
                """, (
                    r["customer_code"], r["customer_name"], r["sale_key_raw"],
                    int(r["year"]), int(r["month"]), float(r["total_amount"]), r["source_sheet"],
                ))
                n_inserted += 1
        conn.commit()
        print(f"บันทึกเรียบร้อย: {n_inserted} แถว")

        # สรุปให้ดูหลัง import
        cur = conn.cursor()
        cur.execute("""
            SELECT year, month, COUNT(*) AS n_customers, SUM(total_amount) AS total
            FROM sales_history_customer_monthly
            GROUP BY year, month ORDER BY year, month
        """)
        print("\nสรุปหลัง import:")
        for row in cur.fetchall():
            print(f"  {row[0]}/{row[1]:02d}  ลูกค้า {row[2]} ราย  ยอดรวม {row[3]:,.2f} บาท")
    except Exception as e:
        conn.rollback()
        print(f"Import failed: {e}")
        raise
    finally:
        app.release_connection(conn)
        app.destroy()


if __name__ == "__main__":
    main()
